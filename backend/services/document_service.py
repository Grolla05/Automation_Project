import os
import logging
from docx import Document
from utils.logger_config import setup_logger

logger = setup_logger()

class DocumentService:
    """
    Serviço para geração de documentos Word (.docx) utilizando layouts pré-definidos.
    """
    
    def __init__(self):
        self.output_dir = 'storage/exports'
        self.layout_dir = 'storage/layout'
        
        # Garante que as pastas existam
        for directory in [self.output_dir, self.layout_dir]:
            if not os.path.exists(directory):
                os.makedirs(directory, exist_ok=True)

    def generate_report(self, file_data, filename="relatorio_ocr.docx", layout_name=None):
        """
        Cria um arquivo Word baseado em um layout ou um documento novo.
        
        Args:
            file_data (list): Lista de dicionários com os dados dos arquivos processados (filename, text, path).
            filename (str): Nome do arquivo de saída.
            layout_name (str): Nome do arquivo de layout (ex: "Bluetooth Low Energy").
        """
        try:
            doc = None
            
            # 1. Tenta carregar o layout se fornecido
            if layout_name:
                # Normaliza o nome para buscar o arquivo .docx
                layout_filename = f"{layout_name}.docx"
                layout_path = os.path.join(self.layout_dir, layout_filename)
                
                if os.path.exists(layout_path):
                    logger.info(f"Usando layout customizado: {layout_path}")
                    doc = Document(layout_path)
                else:
                    logger.warning(f"Layout '{layout_filename}' não encontrado em {self.layout_dir}. Criando documento padrão.")
            
            # 2. Se não houver layout, cria um documento do zero
            if doc is None:
                doc = Document()
                doc.add_heading('Relatório Automático de OCR', 0)
            
            from docx.shared import Inches
            from docx.enum.text import WD_COLOR_INDEX
            import re
            from utils.document_tags import get_text_tags, get_image_tags
            
            if not file_data:
                file_data = [{"filename": "Nenhum arquivo", "text": "Sem dados", "path": None}]
            
            normalized_layout = layout_name.replace('_layout', '') if layout_name else None
            
            tags_texto = {}
            tags_imagem_list = []
            
            # Ordenamos os arquivos alfabeticamente para que: image_teste2.1 seja 1, image_teste2.2 seja 2
            file_data_sorted = sorted(file_data, key=lambda x: str(x.get("filename", "")).lower())
            
            for idx, arquivo in enumerate(file_data_sorted):
                # O sufixo numérico (ex: 1 ou 2)
                sufixo = str(idx + 1)
                
                nome_arq = arquivo.get("filename", "")
                text_ext = arquivo.get("text", "")
                caminho = arquivo.get("path", None)
                
                # Gera as Tags COM SUFIXO e injeta no dicionário gigante
                tags_texto.update(get_text_tags(nome_arq, text_ext, normalized_layout, sufixo=sufixo))
                
                # Para manter compatibilidade retroativa com os outros templates de uma imagem apenas:
                if idx == 0:
                    tags_texto.update(get_text_tags(nome_arq, text_ext, normalized_layout, sufixo=""))
                    
                # E associa a imagem física àquela tag pra hora da injeção
                for timg in get_image_tags(sufixo=sufixo):
                    tags_imagem_list.append({"tag": timg, "caminho": caminho})
                
                if idx == 0:
                    for timg in get_image_tags(sufixo=""):
                        tags_imagem_list.append({"tag": timg, "caminho": caminho})
            
            # Regex para dividir a string preservando as tags (ex: 'Algo [TAG] a mais' -> ['Algo ', '[TAG]', ' a mais'])
            padrao_tags = re.compile('(' + '|'.join(map(re.escape, tags_texto.keys())) + ')')

            def replace_in_paragraphs(paragraphs):
                for paragraph in paragraphs:
                    # Identifica se há tag de imagem
                    for img_config in tags_imagem_list:
                        tag_img = img_config["tag"]
                        caminho_imagem_tg = img_config["caminho"]
                        
                        if tag_img in paragraph.text:
                            paragraph.text = paragraph.text.replace(tag_img, '')
                            # Adiciona a imagem no parágrafo
                            if caminho_imagem_tg and os.path.exists(caminho_imagem_tg) and caminho_imagem_tg.lower().endswith(('.png', '.jpg', '.jpeg')):
                                run = paragraph.add_run()
                                try:
                                    run.add_picture(caminho_imagem_tg, width=Inches(5))
                                except Exception as e:
                                    logger.error(f"Erro ao inserir imagem {caminho_imagem_tg}: {str(e)}")

                    # Substituição de textos com Grifo Amarelo
                    if any(t in paragraph.text for t in tags_texto.keys()):
                        texto_original = paragraph.text
                        paragraph.clear() # Limpa os runs originais do parágrafo
                        
                        partes = padrao_tags.split(texto_original)
                        for parte in partes:
                            if parte in tags_texto:
                                # É uma das nossas Tags! Vamos injetar o dado e grifar de amarelo
                                run_dados = paragraph.add_run(tags_texto[parte])
                                run_dados.font.highlight_color = WD_COLOR_INDEX.YELLOW
                            elif parte:
                                # Texto normal em volta das tags, apenas adicionamos de volta
                                paragraph.add_run(parte)

            # Processa parágrafos soltos
            replace_in_paragraphs(doc.paragraphs)
            
            # Processa parágrafos dentro de tabelas
            for table in doc.tables:
                for row in table.rows:
                    for cell in row.cells:
                        replace_in_paragraphs(cell.paragraphs)
            
            # 4. Salva o resultado
            output_path = os.path.join(self.output_dir, filename)
            doc.save(output_path)
            
            logger.info(f"Documento gerado com sucesso em: {output_path}")
            return output_path
            
        except Exception as e:
            logger.error(f"Erro ao gerar documento Word: {str(e)}")
            raise
